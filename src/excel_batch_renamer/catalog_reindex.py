"""先编排文件目录与图纸目录，再统一重命名工作表、文件夹和 JPG。"""

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Dict, List, Sequence, Tuple
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from excel_batch_renamer.batch_rename_images import (
    batch_rename_images,
)
from excel_batch_renamer.core.page_ranges import build_page_title_map
from excel_batch_renamer.core.naming import extract_folder_sequence
from excel_batch_renamer.infrastructure.pdf_writer import validate_jpeg
from excel_batch_renamer.infrastructure.xlsx_reader import read_image_task


FILE_CATALOG_LABEL = "文件目录"
DRAWING_CATALOG_LABEL = "图纸目录"
MAPPING_WORKBOOK_NAME = "档号与新序号对应表.xlsx"


@dataclass(frozen=True)
class CatalogMappingItem:
    """一卷档案从原档号到全局新序号及目标文件夹的映射。"""

    category: str
    workbook_path: Path
    original_worksheet_name: str
    archive_code: str
    new_sequence: int
    folder_name: str
    source_folder: Path
    target_folder: Path

    @property
    def new_worksheet_name(self) -> str:
        return "{:03d}".format(self.new_sequence)


@dataclass(frozen=True)
class CatalogReindexPlan:
    """两个目录工作簿及全部文件夹在执行前形成的完整编排计划。"""

    file_workbook_path: Path
    drawing_workbook_path: Path
    archive_name_workbook_path: Path
    parent_directory: Path
    mapping_path: Path
    items: Tuple[CatalogMappingItem, ...]


@dataclass(frozen=True)
class CatalogReindexResult:
    """工作表、文件夹、图片和 PDF 的执行汇总。"""

    worksheets: int
    folders_renamed: int
    folders_unchanged: int
    images: int
    images_renamed: int
    images_unchanged: int
    generated_pdfs: int
    mapping_path: Path


def organize_catalog_images(
    file_workbook_path: Path,
    drawing_workbook_path: Path,
    archive_name_workbook_path: Path,
    parent_directory: Path,
) -> CatalogReindexResult:
    """文件目录在前、图纸目录在后，统一编排并执行全部改名与 PDF。"""

    plan = build_catalog_reindex_plan(
        file_workbook_path,
        drawing_workbook_path,
        archive_name_workbook_path,
        parent_directory,
    )
    _write_mapping_workbook(plan)
    _rename_workbook_sheets(plan)
    renamed, unchanged = _rename_folders(plan.items)
    for item in plan.items:
        _renumber_folder_jpgs(item.target_folder)

    file_result = batch_rename_images(
        plan.file_workbook_path,
        plan.parent_directory,
    )
    drawing_result = batch_rename_images(
        plan.drawing_workbook_path,
        plan.parent_directory,
    )
    return CatalogReindexResult(
        worksheets=len(plan.items),
        folders_renamed=renamed,
        folders_unchanged=unchanged,
        images=file_result.total + drawing_result.total,
        images_renamed=file_result.renamed + drawing_result.renamed,
        images_unchanged=file_result.unchanged + drawing_result.unchanged,
        generated_pdfs=(
            file_result.generated_pdfs + drawing_result.generated_pdfs
        ),
        mapping_path=plan.mapping_path,
    )


def build_catalog_reindex_plan(
    file_workbook_path: Path,
    drawing_workbook_path: Path,
    archive_name_workbook_path: Path,
    parent_directory: Path,
) -> CatalogReindexPlan:
    """按档号自然顺序分配全局新序号并完成预检，不依赖标签位置。"""

    file_path = Path(file_workbook_path)
    drawing_path = Path(drawing_workbook_path)
    archive_name_path = Path(archive_name_workbook_path)
    parent = Path(parent_directory)
    if file_path.resolve() == drawing_path.resolve():
        raise ValueError("文件目录和图纸目录必须选择两个不同的 Excel")
    _validate_workbook_name(file_path, FILE_CATALOG_LABEL)
    _validate_workbook_name(drawing_path, DRAWING_CATALOG_LABEL)
    archive_names = _read_archive_names(archive_name_path)
    if not parent.is_dir():
        raise ValueError("图片任务父目录不存在或不是文件夹：{}".format(parent))

    workbook_sheets = (
        (
            FILE_CATALOG_LABEL,
            file_path,
            sorted(_worksheet_names(file_path), key=_archive_code_sort_key),
        ),
        (
            DRAWING_CATALOG_LABEL,
            drawing_path,
            sorted(_worksheet_names(drawing_path), key=_archive_code_sort_key),
        ),
    )
    all_names = [name for _, _, names in workbook_sheets for name in names]
    if not all_names:
        raise ValueError("两个目录工作簿中没有可编排的工作表")
    if len({name.casefold() for name in all_names}) != len(all_names):
        raise ValueError("文件目录与图纸目录包含重复档号")
    if all(_is_new_worksheet_name(name) for name in all_names):
        return _build_repeat_plan(
            file_path,
            drawing_path,
            archive_name_path,
            archive_names,
            parent,
            workbook_sheets,
        )
    if any(_is_new_worksheet_name(name) for name in all_names):
        raise ValueError("两个目录工作簿同时包含原档号和新序号工作表")

    directories = [path for path in parent.iterdir() if path.is_dir()]
    items: List[CatalogMappingItem] = []
    new_sequence = 1
    claimed_folders = set()
    for category, workbook_path, worksheet_names in workbook_sheets:
        for worksheet_name in worksheet_names:
            archive_code = str(worksheet_name).strip()
            folder_name = archive_names.get(archive_code.casefold())
            if folder_name is None:
                raise ValueError(
                    "档号名称 Excel 缺少档号：{}".format(archive_code)
                )
            source_folder, _ = _match_archive_folder(
                directories,
                archive_code,
                claimed_folders,
            )
            target = parent / "{:03d}-{}".format(new_sequence, folder_name)
            _validate_image_inputs(
                workbook_path,
                worksheet_name,
                source_folder,
            )
            items.append(
                CatalogMappingItem(
                    category=category,
                    workbook_path=workbook_path,
                    original_worksheet_name=worksheet_name,
                    archive_code=archive_code,
                    new_sequence=new_sequence,
                    folder_name=folder_name,
                    source_folder=source_folder,
                    target_folder=target,
                )
            )
            claimed_folders.add(source_folder.resolve())
            new_sequence += 1

    _validate_folder_targets(items)
    return CatalogReindexPlan(
        file_workbook_path=file_path,
        drawing_workbook_path=drawing_path,
        archive_name_workbook_path=archive_name_path,
        parent_directory=parent,
        mapping_path=parent / MAPPING_WORKBOOK_NAME,
        items=tuple(items),
    )


def _validate_workbook_name(path: Path, expected_text: str) -> None:
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ValueError("{} Excel 不存在或格式不正确：{}".format(expected_text, path))
    if expected_text not in path.stem:
        raise ValueError("所选 Excel 文件名必须包含“{}”：{}".format(expected_text, path.name))


def _worksheet_names(path: Path) -> List[str]:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return [str(name).strip() for name in workbook.sheetnames]
    finally:
        workbook.close()


def _read_archive_names(path: Path) -> Dict[str, str]:
    """读取仅含“档号、文件夹名称”的对应表，输入表不接受序号列。"""

    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ValueError("档号名称 Excel 不存在或格式不正确：{}".format(path))
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("档号名称 Excel 必须且只能包含一个工作表")
        worksheet = workbook.worksheets[0]
        header = [
            "" if cell.value is None else "".join(str(cell.value).split())
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
        ]
        if "序号" in header:
            raise ValueError("档号名称 Excel 不应包含“序号”列")
        required = ("档号", "文件夹名称")
        indexes = {}
        for name in required:
            matches = [index for index, value in enumerate(header) if value == name]
            if len(matches) != 1:
                raise ValueError("档号名称 Excel 第一行必须包含唯一列：{}".format(name))
            indexes[name] = matches[0]
        result: Dict[str, str] = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if all(value is None or not str(value).strip() for value in row):
                break
            archive_code = str(row[indexes["档号"]]).strip()
            folder_name = str(row[indexes["文件夹名称"]]).strip()
            if not archive_code or archive_code == "None":
                raise ValueError("档号名称 Excel 的档号不能为空")
            if not folder_name or folder_name == "None":
                raise ValueError("档号 {} 的文件夹名称不能为空".format(archive_code))
            key = archive_code.casefold()
            if key in result:
                raise ValueError("档号名称 Excel 包含重复档号：{}".format(archive_code))
            result[key] = folder_name
        if not result:
            raise ValueError("档号名称 Excel 没有有效数据")
        return result
    finally:
        workbook.close()


def _archive_code_sort_key(value: str):
    """把档号中的数字段按数值排序，文本段按大小写不敏感顺序排序。"""

    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.findall(r"\d+|\D+", str(value).strip())
    )


def _is_new_worksheet_name(name: str) -> bool:
    return len(name) == 3 and name.isdigit() and int(name) > 0


def _match_archive_folder(
    directories: Sequence[Path],
    archive_code: str,
    claimed_folders: set,
) -> Tuple[Path, str]:
    matches = []
    for path in directories:
        if path.resolve() in claimed_folders:
            continue
        name_key = path.name.casefold()
        code_key = archive_code.casefold()
        if name_key == code_key:
            matches.append((path, ""))
            continue
        for separator in ("——", "-"):
            prefix = code_key + separator
            if name_key.startswith(prefix):
                matches.append((path, path.name[len(archive_code) + len(separator):]))
                break
    if not matches:
        final_part = archive_code.rsplit("-", 1)[-1]
        if final_part.isdigit() and int(final_part) > 0:
            sequence = int(final_part)
            for path in directories:
                if path.resolve() in claimed_folders:
                    continue
                try:
                    if extract_folder_sequence(path.name) != sequence:
                        continue
                except ValueError:
                    continue
                numeric_prefix = "{:03d}".format(sequence)
                if path.name.startswith(numeric_prefix + "——"):
                    folder_name = path.name[len(numeric_prefix + "——"):]
                elif path.name.startswith(numeric_prefix + "-"):
                    folder_name = path.name[len(numeric_prefix + "-"):]
                else:
                    continue
                matches.append((path, folder_name))
    if not matches:
        raise ValueError("档号 {} 未找到对应直属文件夹".format(archive_code))
    if len(matches) > 1:
        raise ValueError(
            "档号 {} 匹配到多个直属文件夹：{}".format(
                archive_code,
                "、".join(path.name for path, _ in matches),
            )
        )
    return matches[0]


def _validate_image_inputs(
    workbook_path: Path,
    worksheet_name: str,
    folder_path: Path,
) -> None:
    page_titles = build_page_title_map(read_image_task(workbook_path, worksheet_name))
    if not page_titles:
        raise ValueError("工作表 {} 没有可执行的图片任务".format(worksheet_name))
    expected = set(range(1, len(page_titles) + 1))
    if set(page_titles) != expected:
        raise ValueError(
            "工作表 {} 的页次必须从 1 连续到 {}".format(
                worksheet_name,
                len(page_titles),
            )
        )
    images = _ordered_jpgs(folder_path)
    if len(images) != len(expected):
        raise ValueError(
            "档号 {} 的 JPG 数量为 {}，工作表页次为 {} 页".format(
                worksheet_name,
                len(images),
                len(expected),
            )
        )
    for image in images:
        validate_jpeg(image)


def _ordered_jpgs(folder_path: Path) -> List[Path]:
    by_page: Dict[int, Path] = {}
    for path in folder_path.iterdir():
        if not path.is_file() or path.suffix.lower() != ".jpg":
            continue
        prefix = path.name[:3]
        if len(prefix) != 3 or not prefix.isdigit() or int(prefix) < 1:
            raise ValueError("JPG 文件名必须以三位正数页码开头：{}".format(path.name))
        page = int(prefix)
        if page in by_page:
            raise ValueError("同一页码匹配到多个 JPG：{}".format(page))
        by_page[page] = path
    return [by_page[page] for page in sorted(by_page)]


def _validate_folder_targets(items: Sequence[CatalogMappingItem]) -> None:
    source_keys = {item.source_folder.resolve() for item in items}
    target_names = set()
    for item in items:
        key = item.target_folder.name.casefold()
        if key in target_names:
            raise ValueError("多个文件夹将使用同一目标名称：{}".format(item.target_folder.name))
        target_names.add(key)
        if item.target_folder.exists() and item.target_folder.resolve() not in source_keys:
            raise ValueError("文件夹目标名称已被占用：{}".format(item.target_folder))


def _write_mapping_workbook(plan: CatalogReindexPlan) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "档号映射"
    worksheet.append(["目录类别", "档号", "新序号", "文件夹名称"])
    for item in plan.items:
        worksheet.append(
            [
                item.category,
                item.archive_code,
                item.new_worksheet_name,
                item.folder_name,
            ]
        )
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 50
    temporary = plan.mapping_path.with_name(
        ".{}-{}.tmp.xlsx".format(plan.mapping_path.stem, uuid4().hex)
    )
    try:
        workbook.save(str(temporary))
        temporary.replace(plan.mapping_path)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()


def _rename_workbook_sheets(plan: CatalogReindexPlan) -> None:
    for workbook_path in (plan.file_workbook_path, plan.drawing_workbook_path):
        items = [item for item in plan.items if item.workbook_path == workbook_path]
        workbook = load_workbook(str(workbook_path))
        temporary = workbook_path.with_name(
            ".{}-{}.tmp.xlsx".format(workbook_path.stem, uuid4().hex)
        )
        try:
            for index, item in enumerate(items, start=1):
                workbook[item.original_worksheet_name].title = "__tmp_{:03d}".format(index)
            for index, item in enumerate(items, start=1):
                workbook["__tmp_{:03d}".format(index)].title = item.new_worksheet_name
            workbook._sheets.sort(key=lambda worksheet: int(worksheet.title))
            workbook.save(str(temporary))
            temporary.replace(workbook_path)
        finally:
            workbook.close()
            if temporary.exists():
                temporary.unlink()


def _rename_folders(items: Sequence[CatalogMappingItem]) -> Tuple[int, int]:
    changed = [item for item in items if item.source_folder != item.target_folder]
    unchanged = len(items) - len(changed)
    temporary_paths = []
    for item in changed:
        temporary = item.source_folder.parent / (
            ".__catalog_reindex_{}.tmp".format(uuid4().hex)
        )
        item.source_folder.rename(temporary)
        temporary_paths.append((temporary, item.target_folder))
    for temporary, target in temporary_paths:
        temporary.rename(target)
    return len(changed), unchanged


def _renumber_folder_jpgs(folder_path: Path) -> None:
    images = _ordered_jpgs(folder_path)
    staged = []
    for image in images:
        temporary = folder_path / ".__page_{}.tmp".format(uuid4().hex)
        image.rename(temporary)
        staged.append(temporary)
    for page, temporary in enumerate(staged, start=1):
        temporary.rename(folder_path / "{:03d}.jpg".format(page))


def _build_repeat_plan(
    file_path: Path,
    drawing_path: Path,
    archive_name_path: Path,
    archive_names: Dict[str, str],
    parent: Path,
    workbook_sheets,
) -> CatalogReindexPlan:
    mapping_path = parent / MAPPING_WORKBOOK_NAME
    if not mapping_path.is_file():
        raise ValueError("工作表已经是新序号，但未找到档号映射表：{}".format(mapping_path))
    workbook = load_workbook(str(mapping_path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    records = {str(row[2]).zfill(3): row for row in rows if row[2] is not None}
    items = []
    path_by_category = {FILE_CATALOG_LABEL: file_path, DRAWING_CATALOG_LABEL: drawing_path}
    for category, workbook_path, names in workbook_sheets:
        for name in names:
            row = records.get(name)
            if row is None or row[0] != category:
                raise ValueError("档号映射表与当前工作表不一致：{}".format(name))
            archive_code = str(row[1]).strip()
            folder_name = archive_names.get(archive_code.casefold())
            if folder_name is None:
                raise ValueError("档号名称 Excel 缺少档号：{}".format(archive_code))
            matches = []
            for path in parent.iterdir():
                if not path.is_dir():
                    continue
                try:
                    if extract_folder_sequence(path.name) == int(name):
                        matches.append(path)
                except ValueError:
                    continue
            if len(matches) != 1:
                raise ValueError("新序号 {} 未找到对应文件夹".format(name))
            source = matches[0]
            target = parent / "{}-{}".format(name, folder_name)
            _validate_image_inputs(workbook_path, name, source)
            items.append(
                CatalogMappingItem(
                    category=category,
                    workbook_path=path_by_category[category],
                    original_worksheet_name=name,
                    archive_code=archive_code,
                    new_sequence=int(name),
                    folder_name=folder_name,
                    source_folder=source,
                    target_folder=target,
                )
            )
    _validate_folder_targets(items)
    return CatalogReindexPlan(
        file_workbook_path=file_path,
        drawing_workbook_path=drawing_path,
        archive_name_workbook_path=archive_name_path,
        parent_directory=parent,
        mapping_path=mapping_path,
        items=tuple(items),
    )
