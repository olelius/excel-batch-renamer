import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader

from excel_batch_renamer.catalog_reindex import organize_catalog_images


class CatalogReindexTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def _workbook(self, name, sheet_names):
        path = self.root / name
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name in sheet_names:
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.append(["目录说明"])
            worksheet.append(["制表说明"])
            worksheet.append(["序号", "文件题名", "页次"])
            worksheet.append([1, "{}资料".format(sheet_name), "001-002"])
        workbook.save(str(path))
        workbook.close()
        return path

    def _folder(self, archive_code):
        folder = self.root / archive_code
        folder.mkdir()
        for page in (10, 20):
            with Image.new("RGB", (40, 50), color=(page, 80, 160)) as image:
                image.save(str(folder / "{:03d}.jpg".format(page)), "JPEG")
        return folder

    def _archive_names(self, archive_codes):
        path = self.root / "档号名称对应表.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["档号", "文件夹名称"])
        for archive_code in archive_codes:
            worksheet.append([archive_code, "{}卷".format(archive_code)])
        workbook.save(str(path))
        workbook.close()
        return path

    def test_file_catalog_precedes_drawing_catalog_and_matches_by_archive_code(self):
        file_workbook = self._workbook(
            "项目文件目录.xlsx",
            ["I42-4-379", "I42-4-377"],
        )
        drawing_workbook = self._workbook(
            "项目图纸目录.xlsx",
            ["I42-4-380", "I42-4-378"],
        )
        for archive_code in (
            "I42-4-377",
            "I42-4-378",
            "I42-4-379",
            "I42-4-380",
        ):
            self._folder(archive_code)
        archive_names = self._archive_names(
            ("I42-4-377", "I42-4-378", "I42-4-379", "I42-4-380")
        )

        result = organize_catalog_images(
            file_workbook,
            drawing_workbook,
            archive_names,
            self.root,
        )

        self.assertEqual(
            (result.worksheets, result.folders_renamed, result.images,
             result.generated_pdfs),
            (4, 4, 8, 4),
        )
        workbook = load_workbook(str(file_workbook), read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["001", "002"])
        finally:
            workbook.close()
        workbook = load_workbook(str(drawing_workbook), read_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["003", "004"])
        finally:
            workbook.close()

        mapping = load_workbook(str(result.mapping_path), read_only=True, data_only=True)
        try:
            values = list(mapping.active.iter_rows(min_row=2, values_only=True))
        finally:
            mapping.close()
        self.assertEqual(
            [(row[0], row[1], row[2]) for row in values],
            [
                ("文件目录", "I42-4-377", "001"),
                ("文件目录", "I42-4-379", "002"),
                ("图纸目录", "I42-4-378", "003"),
                ("图纸目录", "I42-4-380", "004"),
            ],
        )
        for sequence, archive_code in enumerate(
            ("I42-4-377", "I42-4-379", "I42-4-378", "I42-4-380"),
            start=1,
        ):
            folder = self.root / "{:03d}-{}卷".format(sequence, archive_code)
            self.assertTrue(folder.is_dir())
            self.assertTrue((folder / "001{}资料.jpg".format(archive_code)).is_file())
            self.assertTrue((folder / "002{}资料.jpg".format(archive_code)).is_file())
            self.assertEqual(
                len(PdfReader(folder / "001{}资料.pdf".format(archive_code)).pages),
                2,
            )


if __name__ == "__main__":
    unittest.main()
